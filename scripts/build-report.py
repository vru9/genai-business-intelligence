"""Build the academic report. Requires python-docx and Pillow; render separately with LibreOffice."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
import math
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs'; OUT.mkdir(exist_ok=True)
im=Image.new('RGB',(1500,900),'white'); draw=ImageDraw.Draw(im)
font_path='/System/Library/Fonts/Supplemental/Arial.ttf'
bold_path='/System/Library/Fonts/Supplemental/Arial Bold.ttf'
def font(size,bold=False):
    try:return ImageFont.truetype(bold_path if bold else font_path,size)
    except OSError:return ImageFont.load_default(size=size)
def txt(x,y,t,size=24,bold=False):draw.text((x,y),t,font=font(size,bold),anchor='mm',fill='#273b2e')
boxes=[(40,50,'Data input','Sample dataset or CSV'),(560,50,'Browser analytics','Validate, filter, aggregate'),(1080,50,'Business interface','Dashboard and scenarios'),(560,370,'Server endpoint','Validate and recompute'),(40,670,'Demo engine','Deterministic text rules'),(1080,670,'Gemini service','Question and metrics only')]
for x,y,title,sub in boxes:
    draw.rounded_rectangle((x,y,x+380,y+160),radius=15,fill='#f3f6f3',outline='#7e8f85',width=2)
    txt(x+190,y+59,title,29,True);txt(x+190,y+112,sub,21)
def arrow(a,b):
    draw.line([a,b],fill='#405c48',width=4);ang=math.atan2(b[1]-a[1],b[0]-a[0]);pts=[b,(b[0]-18*math.cos(ang-.45),b[1]-18*math.sin(ang-.45)),(b[0]-18*math.cos(ang+.45),b[1]-18*math.sin(ang+.45))];draw.polygon(pts,fill='#405c48')
arrow((420,130),(550,130));arrow((940,130),(1070,130));arrow((750,210),(750,360));txt(960,285,'Ask a question',21)
arrow((650,540),(300,660));txt(380,550,'No API key',21)
arrow((850,540),(1200,660));txt(1120,550,'API key set',21)
txt(750,875,'Responses return to the interface as plain text',23)
im.save(OUT/'architecture.png')

doc=Document();sec=doc.sections[0];sec.page_height=Inches(11.7);sec.page_width=Inches(8.3);sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.72);sec.left_margin=Inches(.85);sec.right_margin=Inches(.85)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(8);normal.paragraph_format.line_spacing=1.15
for name,size in [('Title',27),('Heading 1',21),('Heading 2',14),('Heading 3',12)]:
    st=doc.styles[name];st.font.size=Pt(size);st.font.bold=name!='Title';st.paragraph_format.space_before=Pt(13);st.paragraph_format.space_after=Pt(9)
    st.paragraph_format.keep_with_next=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Insight BI  |  Group 8  |  Division B  |  Batch B3  |  ');r.font.size=Pt(9);r.font.color.rgb=RGBColor(90,90,90)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='GenAI Powered Business Intelligence and Strategic Insight Generation System'
doc.core_properties.author='Vruddhi Mule; Rajdeep Pandey; Sagar Jadhav'
doc.core_properties.subject='Group 8 Division B Batch B3 college mini project report for academic year 2026–2027'

def p(text='',bold=False):
    para=doc.add_paragraph();r=para.add_run(text);r.bold=bold;return para

def h(text,level=1):doc.add_heading(text,level)
def page(title):doc.add_page_break();h(title)
def bullet(text):
    para=doc.add_paragraph(text,style='List Bullet');para.paragraph_format.space_after=Pt(5);return para

def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,text in enumerate(headers):t.rows[0].cells[i].text=text
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for i,text in enumerate(row):pass
        cells=t.add_row().cells
        for i,text in enumerate(row):cells[i].text=str(text)
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'95');e.set(qn('w:type'),'dxa');marg.append(e)
            pr.append(marg);borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'33465A' if ri==0 else ('F3F5F7' if ri%2==0 else 'FFFFFF'));pr.append(sh)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.05
                for run in para.runs:run.font.size=Pt(10);run.bold=ri==0;run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
    p('')
    return t

def picture(path,width=6.55,caption=''):
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.add_run().add_picture(str(path),width=Inches(width))
    if caption:
        cap=p(caption);cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for r in cap.runs:r.italic=True;r.font.size=Pt(9)

# 1 Cover
p('K.J Somaiya School of Engineering',True).alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Department of IT (AI & DS)').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Division B  |  Batch B3  |  Group 8').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Academic year 2026–2027').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('');p('')
title=doc.add_paragraph('GenAI Powered Business Intelligence and Strategic Insight Generation System',style='Title');title.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Insight BI Application Project Report for Group 8',True).alignment=WD_ALIGN_PARAGRAPH.CENTER
p('');p('')
p('Submitted by Group 8',True).alignment=WD_ALIGN_PARAGRAPH.CENTER
table(['Team member','Roll number'],[['Vruddhi Mule','16014223099'],['Rajdeep Pandey','16014223064'],['Sagar Jadhav','16014223070']],[3.7,2.85])
p('Mini project in generative artificial intelligence and business analytics').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('October 2026').alignment=WD_ALIGN_PARAGRAPH.CENTER

# 2
page('Abstract and report overview')
h('Abstract',2)
p('Insight BI is a web application that transforms structured sales data into business intelligence and strategic explanations. The project addresses the effort required to move from spreadsheet records to an understandable view of revenue, cost, profitability, and performance across business segments. Users can load a sample dataset or upload a CSV, examine a dashboard, ask questions about the selected records, and evaluate business scenarios.')
p('The system combines deterministic numerical analysis with an optional generative AI layer. A shared TypeScript module computes all business metrics. The server validates the selected records again before sending calculated summaries and a question to the Gemini API. This approach grounds generated explanations in explicit numerical evidence. Without a configured API key, the application uses clearly identified rule-based demo responses, allowing the complete demonstration workflow to run without external model access.')
p('The implemented prototype includes CSV validation, regional and date filtering, visual summaries, natural-language questions, scenario calculations, and report exports. Fourteen automated tests verify the core calculations and provider contract. Browser checks verify the main user interactions. Live generative output quality has not been evaluated because no provider key was configured during verification. The application is an educational decision-support prototype, not a validated forecasting or autonomous decision-making system.')
h('Report contents',2)
table(['Section','Page'],[['Introduction and objectives','3'],['Requirements and data model','4'],['System architecture','5'],['Analytical methodology','6'],['Implementation and AI integration','7'],['Application dashboard','8'],['Insight generation interface','9'],['Strategy planning interface','10'],['Testing and evaluation','11'],['Deployment and limitations','12'],['Conclusion and references','13']],[5.5,1.05])

# 3
page('Introduction and objectives')
h('Project context',2)
p('Business data often arrives as rows of dates, categories, locations, revenues, costs, and orders. A spreadsheet can hold these values, but interpretation still requires consistent formulas, segment comparisons, and clear communication. Non-specialist users may struggle to connect a changing revenue chart with a practical next step. The project provides a single workflow for exploring data, explaining observations, and testing assumptions.')
h('Problem statement',2)
p('Develop an accessible application that accepts structured business data, computes reproducible performance indicators, explains relevant patterns in natural language, and helps users compare simple strategic scenarios. The interface must distinguish measured facts from suggestions and distinguish actual language-model responses from a rule-based demonstration.')
h('Objectives',2)
for t in ['Validate CSV inputs before analysis and preserve the previous dataset when an upload is rejected.','Calculate revenue, recorded cost, gross profit, gross margin, orders, and average order value consistently.','Visualize monthly performance and category contributions with region and date filters.','Generate explanations from calculated summaries through an optional Gemini integration.','Provide a rule-based demonstration path when no provider key is configured.','Compare price, volume, and unit-cost assumptions without presenting them as predictions.','Make the application reproducible through documented setup, tests, source code, and an academic report.']:bullet(t)
h('Scope',2)
p('The prototype analyzes a single dataset in Indian rupees and supports up to 10,000 records. It does not include transactional databases, automatic enterprise connectors, user-created accounts, historical report storage, statistical forecasting, or a custom trained model. These boundaries keep the mini project focused on a complete and explainable analytical workflow.')

# 4
page('Requirements and data model')
h('Software and execution requirements',2)
p('Local execution requires Node.js 22.13 or newer, npm, and a modern browser. React and TypeScript implement the interface; Vinext and Vite provide the application build and development server. Server routes run in a Cloudflare Workers-compatible environment. Gemini access is optional and requires an API key and a supported model in the provider account. No application database is required.')
h('CSV schema',2)
table(['Field','Type','Validation or meaning'],[['date','Date text','A real YYYY-MM-DD date between 1900 and 2100'],['category','Text','Required business category, 1–80 characters'],['region','Text','Required geographic segment, 1–80 characters'],['channel','Text','Optional column; absent values default to Direct'],['revenue','Number','Non-negative amount in Indian rupees'],['cost','Number','Non-negative cost using a consistent accounting definition'],['orders','Integer','Non-negative order count summed across records']],[1.05,1,4.5])
h('Data quality controls',2)
p('The parser supports quoted fields, embedded commas, escaped quotation marks, UTF-8 byte-order marks, and common line endings. It rejects missing or duplicate headers, inconsistent column counts, malformed quoting, invalid dates, negative numeric values, blank required numbers, and fractional order counts. Uploads are limited to 2 MB and 10,000 rows. Each grouping dimension is limited to 50 distinct values; date coverage is limited to 3,660 days. Each numeric value is capped at one trillion.')
p('An entire invalid upload is rejected rather than silently dropping rows. Duplicate rows are retained because identical aggregate records may be intentional. Users are responsible for ensuring that repeated records do not double-count business activity. A row may represent one transaction or a sales aggregate, so the application uses the orders column rather than equating row count with order count.')
h('Demonstration dataset',2)
p('The bundled dataset is generated deterministically and contains 288 synthetic records for January–June 2025. It covers four product categories, four Indian regional labels, and three sales channels. It is designed for demonstrations and testing and does not describe an actual business.')

# 5
page('System architecture')
p('The system separates data presentation, deterministic analytics, and optional language generation. A shared analytics module is used in both the browser and the server so the numerical definitions remain consistent.')
picture(OUT/'architecture.png',caption='Figure 1  Data flow and model selection in Insight BI')
h('Browser layer',2)
p('React holds the current dataset, selection filters, search text, and scenario inputs in memory. The analytics module derives filtered totals and grouped data. The dashboard and strategy lab update from that same selection. A refresh restores the sample dataset, which avoids leaving persistent uploaded business records on the device.')
h('Application server layer',2)
p('GET /api/insights reports whether the provider key is configured. POST /api/insights accepts a question and the selected records. It checks the request size and origin, validates the payload, and recalculates the summary. It then selects the rule-based engine or Gemini integration. The API returns plain text with an explicit mode label.')
h('External model boundary',2)
p('When Gemini is enabled, only the question and calculated summaries are transmitted to the provider. Individual sales rows are not sent to the model. Aggregated labels and business values can still be sensitive, so this behavior is disclosed in the interface. API credentials remain in server environment variables.')

# 6
page('Analytical methodology')
h('Business performance metrics',2)
p('For selected records, let R be the sum of revenue, C the sum of recorded cost, and O the sum of orders. The application uses the following definitions.')
table(['Metric','Formula'],[['Gross profit','R − C'],['Gross margin','100 × (R − C) / R'],['Average order value','R / O'],['Category revenue share','100 × category revenue / R'],['Latest month growth','100 × (latest revenue − prior revenue) / prior revenue']],[2.5,4.05])
p('Zero revenue and order denominators are displayed as zero to keep the interface finite; these ratios should be treated as unavailable rather than as evidence of normal business performance. Growth is unavailable when the previous calendar month is absent or has zero revenue. The application does not compare nonconsecutive months as a month-on-month growth rate. A partial month may still be present, and this limitation is shown near the chart and in the insights.')
h('Strategic scenario calculation',2)
p('Let p, v, and c be price, volume, and unit-cost changes expressed as fractions. The scenario assumes that all recorded costs vary with volume. The calculated scenario values are:')
for t in ['Scenario revenue = R × (1 + p) × (1 + v)','Scenario cost = C × (1 + v) × (1 + c)','Scenario gross profit = scenario revenue − scenario cost','Profit change = scenario gross profit − current gross profit']:p(t,True)
p('For a small verification example with revenue 250 and cost 160, a 10% price increase, 20% volume increase, and 10% unit-cost decrease produce revenue 330, cost 172.80, and gross profit 157.20. This calculation is verified by an automated test. Resetting all assumptions to zero reproduces the original profit of 90.')
h('Rule-based observations',2)
p('Insight cards identify the highest-revenue category, the category with the lowest gross margin, and the latest month comparison. Recommended actions are presented as checks or experiments, such as reviewing costs or validating demand. The application does not infer causal reasons from aggregate correlations.')

# 7
page('Implementation and AI integration')
h('Application modules',2)
table(['Module','Implementation responsibility'],[['app/page.tsx','Navigation, dashboard, data explorer, insight questions, strategy controls, and project guide'],['lib/analytics.ts','CSV parser, validation, synthetic sample generator, filtering, grouping, scenario math, and demo answers'],['lib/ai.ts','Provider request construction, model selection, response parsing, and error handling'],['app/api/insights/route.ts','HTTP boundary, server environment configuration, request checks, and responses'],['app/globals.css','Responsive layout, visual states, keyboard focus, and reduced-motion styling']],[2.15,4.4])
h('Prompt construction and grounding',2)
p('The backend builds a JSON context containing the question and computed summary. A system instruction requires the model to use supplied metrics, identify relevant periods and metric names, avoid invented values or causal claims, distinguish observations from hypotheses, and acknowledge missing information. Currency is INR, and profit is explicitly defined as gross profit from the cost column. Temperature is set to 0.25 to encourage restrained responses; this does not guarantee factual correctness.')
p('Dataset labels and questions are designated as untrusted data in the prompt. No model-generated SQL, JavaScript, or HTML is executed. Returned text is rendered by React as text. There is no retrieval database or embedding index; the grounding context consists of numerical aggregates, not retrieved documents.')
h('Modes and failure handling',2)
p('Without GEMINI_API_KEY, the server uses deterministic keyword routes for categories, regions, channels, margins, trends, and summary questions. Unsupported demo questions explain the limitation. With a key, the server calls the Gemini Generate Content endpoint using GEMINI_MODEL. Provider errors, empty answers, and timeouts produce visible errors rather than being silently relabeled as successful AI output. Requests have a 30-second provider timeout and a 40-second browser timeout.')
h('Data handling',2)
p('Uploaded rows remain in browser memory until an insight question is submitted, at which point the selected records reach this application server for validation. The server forwards only aggregates to Gemini. Credentials are excluded from the frontend and ignored by Git. CSV exports protect formula-like text labels. The public demo uses rule-based responses without a provider key. Same-origin checks are implemented; production-grade quotas and per-user authorization remain outside the prototype.')

# 8
page('Application dashboard')
p('The overview combines performance cards, a monthly revenue and gross-profit chart, category contribution, and three rule-based insight cards. Region and date selections affect the same dataset used by the insight service.')
picture(OUT/'screenshots/dashboard.jpg',width=6.55,caption='Figure 2  Desktop dashboard using the complete synthetic sample dataset')
p('The sample produces revenue of INR 4,941,991, recorded cost of INR 3,559,104, gross profit of INR 1,382,887, and 4,026 orders. The gross margin is approximately 27.98%, and average order value is approximately INR 1,227.52. Electronics contributes INR 2,043,888 of revenue. These values are calculated from the bundled sample, not from live business records.')

# 9
page('Insight generation interface')
p('The insight studio provides suggested questions, a free-text input, and the current dataset context. The response label tells the user whether the answer came from Gemini or the rule-based demonstration. Changing the dataset or filters clears a previous answer so it is not presented as belonging to a new selection.')
picture(OUT/'screenshots/ai-insights.jpg',width=6.55,caption='Figure 3  A completed rule-based summary in the insight studio')
p('The demonstrated response reports the sample revenue, profit, and order totals, then describes category contribution, margin review, and monthly revenue change. The displayed demo-mode notice is intentional. The project implements a real provider integration, but does not claim that this screenshot shows a live language-model response.')

# 10
page('Strategy planning interface')
p('The strategy lab lets users change average selling price, sales volume, and cost per unit. Results are recalculated immediately. The panel shows scenario revenue, cost, gross profit, margin, and the difference from the selected baseline. Exporting the report records both assumptions and results.')
picture(OUT/'screenshots/strategy-lab.jpg',width=6.55,caption='Figure 4  Strategy lab with a one percent price increase and unchanged volume and unit cost')
p('For the sample baseline, increasing only price by 1% gives revenue of INR 4,991,410.91 and gross profit of INR 1,432,306.91. The cost stays INR 3,559,104 because the volume and unit-cost assumptions are unchanged. The interface rounds currency to whole rupees. Resetting the controls returns the original values.')
h('Interpretation',2)
p('The output is a conditional calculation. It does not establish that customers will accept a price increase or that order volume will remain unchanged. Fixed costs, capacity limits, demand elasticity, and competitive responses are not included. The results should guide further investigation rather than be treated as a forecast.')

# 11
page('Testing and evaluation')
p('The project includes 14 automated tests implemented with the Node.js test runner. The tests use fixed datasets and explicit expected values so that a calculation error produces a reproducible failure. The provider integration is tested using mock HTTP responses. The automated suite and TypeScript checks passed locally, and the production build completed.')
table(['Test group','Coverage','Result'],[['Core metrics','Revenue, cost, profit, margin, AOV, ranking, and month growth','Pass'],['Selection filters','Region and calendar-month filtering relative to dataset dates','Pass'],['CSV input','Round trip, quoted commas, escaped quotes, missing/duplicate headers, malformed quotes','Pass'],['Validation','Impossible dates, blanks, negatives, fractional orders, row and category limits','Pass'],['Edge cases','Zero denominators, empty summaries, and missing calendar months','Pass'],['Scenarios','Known arithmetic and exact restoration of baseline assumptions','Pass'],['Export handling','Formula-like labels escaped in CSV output','Pass'],['AI contract','Grounded demo text, aggregate-only provider payload, invalid requests, provider errors, empty output','Pass']],[1.3,4.5,.75])
h('Browser verification',2)
p('Manual browser checks confirmed that selecting West updates revenue to INR 1,520,502 across 72 sample records; the CSV sample uploads with 288 valid records; searching Furniture finds 72 records; the summary question returns a labeled demo response; scenario adjustment and reset update the values; and a business report downloads. Desktop and mobile layouts were inspected. A server/client hydration mismatch found during development was corrected before the final verification.')
h('Evaluation limits',2)
p('The tests establish selected numerical and interaction behaviors. They do not establish real-world business usefulness, statistical forecasting accuracy, or language-model reliability. No live Gemini call, user study, large-scale load test, or production security audit was performed. An appropriate next evaluation would use a separately checked business dataset and human ratings of factual support, completeness, and actionability of generated answers.')

# 12
page('Deployment and limitations')
h('Local installation',2)
for t in ['Install Node.js 22.13 or newer and npm.','Clone https://github.com/vru9/genai-business-intelligence and enter the project directory.','Run npm ci, followed by npm run dev. Open the local URL printed in the terminal.','For optional generative mode, copy .dev.vars.example to .dev.vars, set GEMINI_API_KEY and GEMINI_MODEL, and restart.','Run npm test and npm run typecheck to verify the project; npm run build creates the deployment build.']:bullet(t)
h('Repository and hosted access',2)
p('The GitHub repository is public so faculty can inspect the source, sample data, tests, screenshots, and this report. The hosted application uses the Sites platform and is public, so anyone with its link can open the demo. Faculty can also run the repository locally. The public demo works without a provider key using clearly labeled rule-based responses. API keys are not included in the repository.')
p('Repository: https://github.com/vru9/genai-business-intelligence\nHosted address: https://insight-bi-vruddhi.rajdeepvp273.chatgpt.site')
h('Current limitations',2)
for t in ['Uploads and answers are session-memory state and are not saved across refreshes.','The schema uses one currency and does not support negative sales or returns.','The recorded cost definition determines the meaning of profit; taxes and omitted expenses are not inferred.','Monthly comparisons may contain incomplete months and are not adjusted for seasonality.','The scenario model treats every recorded cost as variable and omits price elasticity and fixed expenses.','Demo answers follow limited keyword rules; actual generative answers depend on an external model, credentials, quota, and network availability.','Prompt grounding reduces avoidable errors but cannot guarantee that a model follows every instruction.']:bullet(t)
h('Suggested demonstration',2)
p('Show the dashboard, apply a regional filter, upload the sample CSV, search a category, ask a summary question, explain the mode label, adjust a scenario, reset it, and export the business report. Finish by showing the test suite and explaining the difference between an observation, a recommendation, and a conditional scenario.')

# 13
page('Conclusion and references')
h('Conclusion',2)
p('Insight BI implements an end-to-end educational workflow for structured business analysis: importing data, validating it, computing metrics, visualizing performance, generating explanations, and exploring strategic assumptions. Its main design choice is to calculate numerical facts before asking a language model to explain them. The same analytics module underpins both the dashboard and the server-side context, making the metrics easier to verify.')
p('The project is suitable for demonstrating practical integration of web development, data analysis, prompt construction, and API-based generative AI. A working rule-based mode makes the application demonstrable without credentials, while explicit labels prevent that mode from being confused with actual generative output. The test suite and documented assumptions support a more transparent presentation of both the capabilities and limits of the prototype.')
h('Future enhancements',2)
for t in ['Add persistent datasets and saved reports with per-user access controls.','Introduce configurable column mapping, multiple currencies, returns handling, and richer data quality diagnostics.','Evaluate live model responses against a human-reviewed set of business questions and expected facts.','Add statistical forecasting only with backtesting, error measures, and uncertainty estimates.','Extend scenarios with fixed and variable cost separation, demand elasticity, and sensitivity analysis.','Add request quotas, audit trails, and cost monitoring before wider production use.']:bullet(t)
h('References',2)
refs=[('Google AI for Developers','Gemini Generate Content API','https://ai.google.dev/api/generate-content'),('React','Learning React and component state','https://react.dev/learn'),('Microsoft','TypeScript Handbook','https://www.typescriptlang.org/docs/handbook/intro.html'),('Cloudflare','Cloudflare Workers documentation','https://developers.cloudflare.com/workers/'),('Node.js','Test runner documentation','https://nodejs.org/api/test.html'),('Project source','Insight BI repository and test suite','https://github.com/vru9/genai-business-intelligence')]
for i,(author,title,url) in enumerate(refs,1):
    para=p(f'{i}. {author}. {title}. {url}')
    for r in para.runs:r.font.size=Pt(9)
p('Report prepared for academic year 2026–2027. Application and test evidence recorded on 16 October 2026.').runs[0].font.size=Pt(9)
for element in doc.styles.element.iter(qn('w:pBdr')):
    element.getparent().remove(element)
for element in doc.element.iter(qn('w:pBdr')):
    element.getparent().remove(element)
doc.save(OUT/'project-report.docx');print(OUT/'project-report.docx')
